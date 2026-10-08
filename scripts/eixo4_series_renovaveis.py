#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Eixo 4 / I4.3.2 (Participação de renováveis, PER) — coletor de séries históricas.

Roda do zero, a partir da raiz:  python scripts/eixo4_series_renovaveis.py
Opções:  --so ons|siga|big|cfurh|gd|testes   (roda só uma parte; padrão: todas)

Contexto. O painel mostra, para o valor atual, um PROXY: % renovável (hídrica, solar,
eólica, biomassa) da POTÊNCIA fiscalizada das usinas em operação (SIGA/ANEEL). Nenhuma
fonte publica essa participação por UF ao longo do tempo, e reconstruir pelo
DatEntradaOperacao do SIGA não funciona (viés de sobrevivência e de datação; medido em
dashboard/public/data/csv/diagnostico_siga_reconstrucao.csv). Este script NÃO refaz essa
reconstrução. Ele junta quatro conjuntos de séries que não sofrem desses vieses:

  1. FOTOGRAFIAS de potência (mesmo conceito do painel), cada uma uma lista de usinas em
     operação NAQUELA data, logo sem viés de sobrevivência:
       a) ANEEL/BIG ("Capacidade de Geração do Brasil", páginas GeracaoTipoFase), guardadas
          pelo Wayback Machine: nov/2002, jan/2008, ago/2011, mar/2016;
       b) ANEEL/SIGA, CSVs guardados pelo Wayback: ago/2022, mar/2024, set/2024, jan/2025,
          mai/2025, set/2025, jan/2026; mais o CSV atual do CKAN.
     Teste que separa isto da reconstrução morta: o total por UF de cada fotografia é
     comparado com a capacidade oficial da ANEEL (capacidade-instalada-por-unidade-da-
     federacao) em dezembro anterior e seguinte (validacao_fotografias_vs_aneel.csv).
  2. ENERGIA (GWh), ONS "geracao-usina-2" (geração por usina, horária, 2000 em diante):
     UF x combustível x ano. Cobre só usinas que o ONS monitora (SIN e, depois da
     interligação, parte dos antigos sistemas isolados): ver colunas de cobertura.
  3. HÍDRICA em GWh por UF, CFURH/ANEEL (compensação financeira, 1997 em diante).
  4. Geração distribuída solar (ANEEL), potência acumulada por UF e ano de conexão.
  5. EPE Anuário (Tabelas 2.1, 2.2, 2.5, 2.8): contexto e conferência de cobertura.
  6. Uma RECONSTRUÇÃO anual de 2008 em diante ENTRE as fotografias (parte_hibrida), com o erro medido por
     leave-one-out. Não usa DatEntradaOperacao: ver o docstring de parte_hibrida.

Observações. (a) O ponto "base do painel" (2026-08-25) só entra se dados/aneel/siga.csv existir (ele é
gerado por scripts/eixo4_aneel_siga.py, que não é deste coletor); sem ele a execução do zero omite esse ponto e o
teste contra valores.csv. (b) O Wayback derruba conexões com frequência: há retry com recuo e as capturas estão
fixadas por timestamp (não depende do CDX). (c) Quase 1 GB de download na primeira execução (ONS ~730 MB).

Saídas em dados/eixo4_series/I4.3.2/ (bruto em .../bruto/, que é cache e pode ser apagado).
"""
import argparse
import csv
import glob
import html
import io
import json
import re
import sys
import time
import unicodedata
from pathlib import Path

import pandas as pd
import requests

sys.stdout.reconfigure(encoding="utf-8")

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "dados" / "eixo4_series" / "I4.3.2"
BRUTO = SAIDA / "bruto"
CODIGO = "I4.3.2"
HDR = {"User-Agent": "Mozilla/5.0 (observatorio-amazonia-2050; coleta de series)"}
UFS = ["AC", "AP", "AM", "MA", "MT", "PA", "RO", "RR", "TO"]
RENOVAVEIS = {"Hídrica", "Eólica", "Solar", "Biomassa"}  # definição do painel (SIGA)
VALORES_CSV = RAIZ / "dashboard" / "conteudo" / "valores.csv"
SIGA_PAINEL = RAIZ / "dados" / "aneel" / "siga.csv"  # base 25/08/2026, a do painel

CKAN_ANEEL = "https://dadosabertos.aneel.gov.br/api/3/action/package_show?id="
CKAN_ONS = "https://dados.ons.org.br/api/3/action/package_show?id="
URL_SIGA_ATUAL = ("https://dadosabertos.aneel.gov.br/dataset/6d90b77c-c5f5-4d81-bdec-7bc619494bb9/resource/"
                  "11ec447d-698d-4ab8-977f-b424d5deee6a/download/siga-empreendimentos-geracao.csv")
URL_CAP_UF = ("https://dadosabertos.aneel.gov.br/dataset/cec20fdd-97e4-40a8-870f-c63339f5d8b7/resource/"
              "6fbee0f8-2617-4879-a69a-6b7892f12dad/download/capacidade-instalada-geracao-uf.csv")
URL_CFURH = ("https://dadosabertos.aneel.gov.br/dataset/855a73ed-f4e0-4e88-8467-c330ba0223d3/resource/"
             "df09e492-0927-4873-8488-31d854a14709/download/cfurh-geracao.csv")
URL_EPE = ("https://www.epe.gov.br/sites-pt/publicacoes-dados-abertos/publicacoes/PublicacoesArquivos/"
           "publicacao-160/topico-168/anuario-workbook.xlsx")
URL_LIBERACAO = ("https://dadosabertos.aneel.gov.br/dataset/2b2ace01-5692-4636-8c99-2a84ff094f4f/resource/"
                 "75419902-c692-498b-a6ef-85f6d4beb5b2/download/"
                 "unidades-geradoras-liberadas-operacao-comercial-detalhado.csv")
NOMES_UF = {"Rondônia": "RO", "Acre": "AC", "Amazonas": "AM", "Roraima": "RR", "Pará": "PA", "Amapá": "AP",
            "Tocantins": "TO", "Maranhão": "MA", "Mato Grosso": "MT"}
URL_GD_FV = ("https://dadosabertos.aneel.gov.br/dataset/5e0fafd2-21b9-4d5b-b622-40438d40aba2/resource/"
             "703c4cb8-b7e2-4f27-a9bb-7e55324a88a4/download/"
             "empreendimento-gd-informacoes-tecnicas-fotovoltaica.parquet")
WB = "https://web.archive.org/web/{ts}id_/{url}"  # id_ = bytes originais, sem a moldura do Wayback

# ---- Capturas do Wayback fixadas (o CDX da Internet Archive cai com frequência) -------------
SIGA_WB = [  # (timestamp, arquivo original). Data de referência vem de dentro do CSV.
    ("20220820015327", "11ec447d-698d-4ab8-977f-b424d5deee6a/download/siga-empreendimentos-geracao.csv"),
    ("20240303090125", "11ec447d-698d-4ab8-977f-b424d5deee6a/download/siga-empreendimentos-geracao.csv"),
    ("20240918222633", "11ec447d-698d-4ab8-977f-b424d5deee6a/download/siga-empreendimentos-geracao.csv"),
    ("20250106221759", "2f65a1b0-19b8-4360-8238-b34ab4693d55/download/siga-empreendimentos-geracao-diario.csv"),
    ("20250531061642", "11ec447d-698d-4ab8-977f-b424d5deee6a/download/siga-empreendimentos-geracao.csv"),
    ("20250926220247", "11ec447d-698d-4ab8-977f-b424d5deee6a/download/siga-empreendimentos-geracao.csv"),
    ("20260117053548", "2f65a1b0-19b8-4360-8238-b34ab4693d55/download/siga-empreendimentos-geracao-diario.csv"),
]
SIGA_WB_BASE = "https://dadosabertos.aneel.gov.br/dataset/6d90b77c-c5f5-4d81-bdec-7bc619494bb9/resource/"
_BIG = "aplicacoes/capacidadebrasil/GeracaoTipoFase.asp?tipo={t}&fase=3"
BIG_WB = {  # fotografia -> {tipo: (timestamp, URL original)}
    "2002-11": {"data_ref": "2002-11-11", "paginas": {
        "1": ("20021111043010", "http://www.aneel.gov.br:80/" + _BIG.format(t=1)),
        "2": ("20021111043700", "http://www.aneel.gov.br:80/" + _BIG.format(t=2)),
        "5": ("20021111044051", "http://www.aneel.gov.br:80/" + _BIG.format(t=5)),
        "7": ("20021111042924", "http://www.aneel.gov.br:80/" + _BIG.format(t=7)),
        "9": ("20021111042745", "http://www.aneel.gov.br:80/" + _BIG.format(t=9))}},
    "2008-01": {"data_ref": "2008-01-12", "paginas": {
        "1": ("20080112221810", "http://www.aneel.gov.br:80/" + _BIG.format(t=1)),
        "2": ("20080112223602", "http://www.aneel.gov.br:80/" + _BIG.format(t=2)),
        "5": ("20080112223616", "http://www.aneel.gov.br:80/" + _BIG.format(t=5)),
        "7": ("20080112223632", "http://www.aneel.gov.br:80/" + _BIG.format(t=7)),
        "8": ("20080112223637", "http://www.aneel.gov.br:80/" + _BIG.format(t=8)),
        "9": ("20080112223642", "http://www.aneel.gov.br:80/" + _BIG.format(t=9)),
        "10": ("20080112223537", "http://www.aneel.gov.br:80/" + _BIG.format(t=10))}},
    "2011-08": {"data_ref": "2011-08-12", "paginas": {
        "1": ("20110812101904", "http://www.aneel.gov.br/" + _BIG.format(t=1)),
        "2": ("20110812101757", "http://www.aneel.gov.br/" + _BIG.format(t=2)),
        "5": ("20110812101440", "http://www.aneel.gov.br/" + _BIG.format(t=5)),
        "7": ("20110812101503", "http://www.aneel.gov.br/" + _BIG.format(t=7)),
        "8": ("20110812101533", "http://www.aneel.gov.br/" + _BIG.format(t=8)),
        "9": ("20110812101619", "http://www.aneel.gov.br/" + _BIG.format(t=9)),
        "10": ("20110812101411", "http://www.aneel.gov.br/" + _BIG.format(t=10))}},
    "2016-03": {"data_ref": "2016-03-15", "paginas": {
        "1": ("20160315034846", "http://www2.aneel.gov.br/" + _BIG.format(t=1)),
        "2": ("20160315152859", "http://www2.aneel.gov.br/" + _BIG.format(t=2)),
        "5": ("20160315153015", "http://www2.aneel.gov.br/" + _BIG.format(t=5)),
        "7": ("20160315030804", "http://www2.aneel.gov.br/" + _BIG.format(t=7)),
        "9": ("20160315154758", "http://www2.aneel.gov.br/" + _BIG.format(t=9)),
        "10": ("20160315154823", "http://www2.aneel.gov.br/" + _BIG.format(t=10)),
        "12": ("20160417092341", "http://www2.aneel.gov.br:80/" + _BIG.format(t=12))}},
}
# tipo da página do BIG -> grupo de fonte (UTE, tipo 2, é aberta pelo combustível)
BIG_TIPO = {"1": "Hídrica", "5": "Hídrica", "10": "Hídrica", "7": "Eólica", "8": "Solar", "12": "Solar",
            "9": "Nuclear", "2": None}
# combustíveis de UTE que a ANEEL classifica como biomassa (mesma lista do SIGA atual)
BIOMASSA = {"Bagaço de Cana de Açúcar", "Biogás - Floresta", "Biogás - RA", "Biogás - RU", "Biogás-AGR",
            "Capim Elefante", "Carvão - RU", "Carvão Vegetal", "Casca de Arroz", "Etanol",
            "Gás de Alto Forno - Biomassa", "Lenha", "Licor Negro", "Resíduos Florestais",
            "Resíduos Sólidos Urbanos - RU", "Óleos vegetais"}


NOTAS_VALIDACAO = {
    ("MA", "2008-01-12"): ("a série oficial da ANEEL contava Boa Esperança (237,3 MW, hídrica) no MA até 2008; "
                           "a CEG atual dela é do PI, convenção usada aqui e no painel. 17,0 + 237,3 = 254,3 MW, "
                           "o valor oficial de dez/2008"),
}


# ------------------------------------------------------------------------------------------
def baixar(url, destino, minimo=1000, tentativas=6, pausa=6, timeout=300):
    """GET com retry e recuo; guarda em cache e não baixa de novo."""
    destino = Path(destino)
    if destino.exists() and destino.stat().st_size >= minimo:
        return destino
    destino.parent.mkdir(parents=True, exist_ok=True)
    ultimo = None
    for i in range(tentativas):
        try:
            r = requests.get(url, headers=HDR, timeout=timeout)
            if r.status_code == 200 and len(r.content) >= minimo:
                destino.write_bytes(r.content)
                return destino
            ultimo = f"HTTP {r.status_code}, {len(r.content)} bytes"
        except requests.RequestException as e:
            ultimo = str(e)[:100]
        print(f"   tentativa {i + 1}/{tentativas} falhou ({ultimo}): {url[-80:]}")
        time.sleep(pausa * (i + 1))
    raise RuntimeError(f"não consegui baixar {url}: {ultimo}")


def ler_texto(caminho):
    raw = Path(caminho).read_bytes()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("latin-1")


def num_br(s):
    s = str(s).strip().replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\(.*?\)", " ", s)
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def gravar(df, nome):
    SAIDA.mkdir(parents=True, exist_ok=True)
    df.to_csv(SAIDA / nome, index=False, encoding="utf-8", quoting=csv.QUOTE_MINIMAL)
    print(f"   -> {nome}: {len(df)} linhas")


def capacidade_oficial():
    """Capacidade instalada por UF em dezembro (ANEEL), MW. Usada só para validar."""
    p = baixar(URL_CAP_UF, BRUTO / "aneel_capacidade_uf.csv")
    d = pd.read_csv(io.StringIO(ler_texto(p)), sep=";", decimal=",")
    d = d[(d.MesReferencia == 12) & d.SigUF.isin(UFS)]
    return (d.pivot(index="SigUF", columns="AnoReferencia", values="MdaPotenciaInstaladakW") / 1000)


def per_linhas(pot, data_ref, fonte_dado, origem_url, extra=None):
    """pot: DataFrame UF x origem (MW). Devolve linhas de PER por UF e AL."""
    linhas = []
    for uf in UFS + ["AL"]:
        if uf == "AL":
            tot = pot.sum().sum()
            ren = pot[[c for c in pot.columns if c in RENOVAVEIS]].sum().sum()
        else:
            if uf not in pot.index:
                continue
            tot = pot.loc[uf].sum()
            ren = pot.loc[uf, [c for c in pot.columns if c in RENOVAVEIS]].sum()
        if tot <= 0:
            continue
        linhas.append({"uf": uf, "ano": int(str(data_ref)[:4]), "data_ref": str(data_ref)[:10],
                       "valor": round(ren / tot * 100, 2), "unidade": "% da potência fiscalizada",
                       "fonte": fonte_dado, "potencia_total_mw": round(tot, 1),
                       "potencia_renovavel_mw": round(ren, 1), "origem_url": origem_url})
    return linhas


# ============================== PARTE 1a: BIG (Wayback) =====================================
def parse_big(caminho):
    t = Path(caminho).read_bytes().decode("latin-1")
    usinas, total, cab = [], None, None
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", t, flags=re.S | re.I):
        cel = [re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", c))).strip()
               for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, flags=re.S | re.I)]
        if not cel:
            continue
        if cel[0] in ("Usina", "CEG"):
            cab = cel
            continue
        if cel[0].startswith("Total:"):
            n = int(re.search(r"Total:\s*([\d.]+)", cel[0]).group(1).replace(".", ""))
            m = re.search(r"([\d.]+(?:,\d+)?)\s*kW", " ".join(cel[1:]))
            total = (n, num_br(m.group(1)) if m else None)
            continue
        if cab and len(cel) == len(cab):
            usinas.append(dict(zip(cab, cel)))
    return usinas, total


def fotografias_big():
    print("\n[1a] ANEEL/BIG no Wayback (fotografias de potência em operação)")
    paginas = {}
    for foto, cfg in BIG_WB.items():
        for tipo, (ts, url) in cfg["paginas"].items():
            dest = BRUTO / "wayback" / f"big_{foto}_tipo{tipo}_{ts}.html"
            baixar(WB.format(ts=ts, url=url), dest, minimo=500, pausa=8)
            paginas[(foto, tipo)] = parse_big(dest)
            time.sleep(1)
    # mapa nome -> UF principal, para usinas interestaduais das fotografias sem CEG
    mapa = {}
    for (foto, tipo), (usinas, _) in paginas.items():
        if foto != "2016-03":
            continue
        for u in usinas:
            m = re.match(r"[A-Z]{3}\.[A-Z]{2}\.([A-Z]{2})\.", u.get("CEG", ""))
            if m:
                mapa[norm(u["Usina"])] = m.group(1)
    if SIGA_PAINEL.exists():
        s = pd.read_csv(io.StringIO(ler_texto(SIGA_PAINEL)), sep=";", dtype=str)
        for nome, uf in zip(s.NomEmpreendimento, s.SigUFPrincipal):
            mapa.setdefault(norm(nome), uf)
    linhas, interestaduais = [], []
    for (foto, tipo), (usinas, total) in paginas.items():
        col = "Potência Fiscalizada (kW)" if usinas and "Potência Fiscalizada (kW)" in usinas[0] else "Potência (kW)"
        soma = sum(num_br(u[col]) or 0 for u in usinas)
        # o total impresso na página é de fiscalizada (2008, 2011) ou de outorgada (2016): vale se bater com uma delas
        somas = [soma] + ([sum(num_br(u["Potência Outorgada (kW)"]) or 0 for u in usinas)]
                          if usinas and "Potência Outorgada (kW)" in usinas[0] else [])
        if total and total[1] and not any(abs(x - total[1]) <= 1 for x in somas):
            print(f"   AVISO {foto} tipo {tipo}: somas {somas} kW != total da página {total[1]:.0f} kW")
        if total and len(usinas) != total[0]:
            print(f"   AVISO {foto} tipo {tipo}: {len(usinas)} usinas lidas != {total[0]} da página")
        for u in usinas:
            m = re.match(r"[A-Z]{3}\.[A-Z]{2}\.([A-Z]{2})\.", u.get("CEG", ""))
            ufs = list(dict.fromkeys(re.findall(r"- ([A-Z]{2})", u["Município"])))
            if m:
                uf = m.group(1)  # a UF da CEG é a "principal" da ANEEL
            elif len(ufs) <= 1:
                uf = ufs[0] if ufs else None
            else:
                uf = mapa.get(norm(u["Usina"]))
                if uf not in ufs:
                    uf = ufs[0]
                if any(x in UFS for x in ufs):
                    interestaduais.append((foto, u["Usina"], "/".join(ufs), uf))
            grupo = BIG_TIPO[tipo]
            if grupo is None:  # UTE: abre pelo combustível
                comb = u.get("Fonte Nível 2") or u.get("Combustível") or ""
                cls = u.get("Classe Combustível")
                grupo = ("Biomassa" if cls == "Biomassa" else "Fóssil") if cls else \
                        ("Biomassa" if comb in BIOMASSA else "Fóssil")
            linhas.append({"foto": foto, "uf": uf, "grupo": grupo, "kw": num_br(u[col]) or 0.0})
    if interestaduais:
        print(f"   usinas interestaduais (UF principal resolvida pela CEG/SIGA): {len(interestaduais)}")
    df = pd.DataFrame(linhas)
    df = df[df.uf.isin(UFS)]
    g = df.groupby(["foto", "uf", "grupo"]).kw.sum().div(1000).round(2).reset_index()
    g["data_ref"] = g.foto.map(lambda f: BIG_WB[f]["data_ref"])
    g = g.rename(columns={"kw": "potencia_mw", "grupo": "origem"})
    g["fonte_dado"] = "ANEEL/BIG via Wayback"
    g["unidade"] = "MW"
    return g[["data_ref", "fonte_dado", "uf", "origem", "potencia_mw", "unidade"]]


# ============================== PARTE 1b: SIGA (Wayback + atual) ============================
def pot_siga(caminho):
    d = pd.read_csv(io.StringIO(ler_texto(caminho)), sep=";", dtype=str)
    ref = d["DatGeracaoConjuntoDados"].iloc[0]
    d = d[d.DscFaseUsina == "Operação"].copy()
    d["kw"] = pd.to_numeric(d.MdaPotenciaFiscalizadaKw.str.replace(".", "", regex=False)
                            .str.replace(",", ".", regex=False), errors="coerce").fillna(0)
    d = d[d.SigUFPrincipal.isin(UFS)]
    g = (d.groupby(["SigUFPrincipal", "DscOrigemCombustivel"]).kw.sum() / 1000).round(2).reset_index()
    g.columns = ["uf", "origem", "potencia_mw"]
    return ref[:10], g


def fotografias_siga():
    print("\n[1b] ANEEL/SIGA: CSVs do Wayback e o CSV atual do CKAN")
    saida = []
    for ts, arq in SIGA_WB:
        url = SIGA_WB_BASE + arq
        dest = BRUTO / "wayback_siga" / f"siga_{ts}.csv"
        baixar(WB.format(ts=ts, url=url), dest, minimo=100000, pausa=8)
        ref, g = pot_siga(dest)
        g["data_ref"], g["fonte_dado"] = ref, "ANEEL/SIGA via Wayback"
        g["url"] = WB.format(ts=ts, url=url)
        saida.append(g)
        time.sleep(1)
    atual = baixar(URL_SIGA_ATUAL, BRUTO / "siga_atual.csv", minimo=100000)
    ref, g = pot_siga(atual)
    g["data_ref"], g["fonte_dado"], g["url"] = ref, "ANEEL/SIGA (CKAN, download desta execução)", URL_SIGA_ATUAL
    saida.append(g)
    if SIGA_PAINEL.exists():  # a base do painel (25/08/2026), para o teste de linha de base
        ref, g = pot_siga(SIGA_PAINEL)
        g["data_ref"], g["fonte_dado"] = ref, "ANEEL/SIGA (base do painel, dados/aneel/siga.csv)"
        g["url"] = URL_SIGA_ATUAL
        saida.append(g)
    return pd.concat(saida, ignore_index=True)


def parte_potencia():
    big = fotografias_big()
    siga = fotografias_siga()
    big["url"] = "web.archive.org (ver BIG_WB no script)"
    todas = pd.concat([big.assign(unidade="MW"), siga.assign(unidade="MW")], ignore_index=True)
    todas = todas.drop_duplicates(["data_ref", "fonte_dado", "uf", "origem"])
    todas = todas.sort_values(["data_ref", "fonte_dado", "uf", "origem"])
    gravar(todas[["data_ref", "fonte_dado", "uf", "origem", "potencia_mw", "unidade", "url"]],
           "fotografias_potencia_uf_fonte.csv")
    linhas = []
    for (ref, fonte), g in todas.groupby(["data_ref", "fonte_dado"]):
        pot = g.pivot_table(index="uf", columns="origem", values="potencia_mw", aggfunc="sum", fill_value=0)
        linhas += per_linhas(pot, ref, fonte, g.url.iloc[0])
    per = pd.DataFrame(linhas).sort_values(["uf", "data_ref", "fonte"])
    # "ano" é o ano da data da captura. As capturas de janeiro-fevereiro mostram o fim do ano anterior (a de
    # 2008-01-12 bate com dez/2007 da ANEEL em 7 das 9 UFs). Para juntar por (uf, ano) sem duplicar linhas:
    # filtrar principal_do_ano == True e usar ano_referencia.
    dt = pd.to_datetime(per.data_ref)
    per["ano_referencia"] = [t.year - 1 if t.month <= 2 else t.year for t in dt]
    dist = [abs((pd.Timestamp(f"{a}-12-31") - t).days) for a, t in zip(per.ano_referencia, dt)]
    per["dias_de_31_dez"] = dist
    per["principal_do_ano"] = per.groupby(["uf", "ano_referencia"]).dias_de_31_dez.transform("min") == per.dias_de_31_dez
    gravar(per, "per_potencia_fotografias_uf.csv")

    # validação: total de cada fotografia contra a capacidade oficial da ANEEL (dez anterior e seguinte)
    of = capacidade_oficial()
    val = []
    for r in per[per.uf != "AL"].itertuples():
        a = int(r.ano)
        ant, seg = a - 1, a
        # fotografia de jan-fev: o estado reflete o fim do ano anterior; as demais ficam entre dez(a-1) e dez(a)
        lo_hi = [of.loc[r.uf, y] for y in (a - 1, a) if y in of.columns]
        if len(lo_hi) < 2:
            continue
        lo, hi = min(lo_hi), max(lo_hi)
        dentro = lo * 0.98 <= r.potencia_total_mw <= hi * 1.02
        val.append({"data_ref": r.data_ref, "fonte": r.fonte, "uf": r.uf,
                    "total_fotografia_mw": r.potencia_total_mw,
                    "aneel_dez_anterior_mw": round(of.loc[r.uf, a - 1], 1),
                    "aneel_dez_mesmo_ano_mw": round(of.loc[r.uf, a], 1),
                    "dentro_do_intervalo_oficial_2pct": bool(dentro),
                    "observacao": NOTAS_VALIDACAO.get((r.uf, r.data_ref), "")})
    val = pd.DataFrame(val)
    gravar(val, "validacao_fotografias_vs_aneel.csv")
    for fonte, g in val.groupby("fonte"):
        print(f"   {fonte}: {g.dentro_do_intervalo_oficial_2pct.sum()}/{len(g)} pontos UF dentro do intervalo oficial (+-2%)")
    return per


# ============================== PARTE EPE: Anuário (tabelas por UF) =========================
def baixar_epe():
    dest = BRUTO / "epe_anuario_workbook.xlsx"
    if dest.exists() and dest.stat().st_size > 100000:
        return dest
    try:
        return baixar(URL_EPE, dest, minimo=100000, tentativas=2)
    except RuntimeError:  # a cadeia de certificados do site da EPE costuma falhar
        import urllib3
        urllib3.disable_warnings()
        r = requests.get(URL_EPE, headers=HDR, timeout=300, verify=False)
        r.raise_for_status()
        dest.write_bytes(r.content)
        return dest


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def parte_epe():
    import openpyxl
    print("\n[EPE] Anuário Estatístico de Energia Elétrica: Tabelas 2.1, 2.2 e 2.5 (UF x ano)")
    wb = openpyxl.load_workbook(baixar_epe(), read_only=True, data_only=True)
    cfg = {"Tabela 2.1": ("capacidade instalada total", "MW"),
           "Tabela 2.2": ("capacidade instalada de MMGD (micro e minigeração distribuída)", "MW"),
           "Tabela 2.5": ("geração elétrica total (todas as fontes)", "GWh")}
    linhas = []
    for aba, (desc, un) in cfg.items():
        anos = None
        for row in wb[aba].iter_rows(values_only=True):
            cel = list(row)
            if anos is None:
                cand = [(i, int(str(c).strip())) for i, c in enumerate(cel)
                        if str(c).strip().isdigit() and len(str(c).strip()) == 4]
                if len(cand) >= 5:
                    anos = cand
                continue
            nome = str(cel[1]).strip() if len(cel) > 1 and cel[1] is not None else ""
            if nome in NOMES_UF:
                for i, ano in anos:
                    v = _num(cel[i]) if i < len(cel) else None
                    if v is not None:
                        linhas.append({"tabela": aba.replace("Tabela ", ""), "descricao": desc, "uf": NOMES_UF[nome],
                                       "ano": ano, "valor": round(v, 2), "unidade": un,
                                       "fonte": "EPE Anuário Estatístico de Energia Elétrica 2026 (ano-base 2025), workbook"})
    df = pd.DataFrame(linhas).sort_values(["tabela", "uf", "ano"])
    gravar(df, "epe_anuario_uf_ano.csv")
    # Tabela 2.8: renovabilidade da oferta interna de ELETRICIDADE, só nacional (contexto)
    br, anos = [], None
    for row in wb["Tabela 2.8"].iter_rows(values_only=True):
        cel = list(row)
        if anos is None:
            cand = [(i, int(str(c).strip())) for i, c in enumerate(cel)
                    if str(c).strip().isdigit() and len(str(c).strip()) == 4]
            if len(cand) >= 5:
                anos = cand
            continue
        nome = str(cel[1]).strip() if len(cel) > 1 and cel[1] is not None else ""
        if nome in ("Renováveis", "Não Renováveis", "Renováveis (%)"):
            for i, ano in anos:
                v = _num(cel[i]) if i < len(cel) else None
                if v is not None:
                    br.append({"uf": "BR", "ano": ano, "serie": nome, "valor": round(v, 2),
                               "unidade": "%" if "%" in nome else "GWh",
                               "fonte": "EPE Anuário 2026, Tabela 2.8 (Oferta Interna de Energia Elétrica; BEN)"})
    gravar(pd.DataFrame(br).sort_values(["serie", "ano"]), "epe_renovabilidade_brasil_ano.csv")
    return df


# ============================== PARTE 1c: reconstrução com âncoras ==========================
def parte_hibrida(epe=None):
    """PER de potência por UF no fim de cada ano, ENTRE fotografias reais (2008 em diante).

    Diferença para a reconstrução morta (DatEntradaOperacao do SIGA): (i) o ponto de partida e o de
    chegada de cada trecho são listas reais de usinas em operação naquela data; (ii) as adições vêm do
    cadastro de LIBERAÇÃO PARA OPERAÇÃO COMERCIAL, que data cada UNIDADE geradora (Belo Monte entra
    unidade a unidade, não numa linha só); (iii) o que sai de operação (térmicas aposentadas) não é
    reconstruído: é o "erro de fechamento" entre as duas fotografias, distribuído linearmente no tempo,
    tanto no numerador (renovável) quanto no total. A qualidade é medida por leave-one-out: prever uma
    fotografia do meio só com as vizinhas (validacao_reconstrucao_leave_one_out.csv).
    O total oficial da ANEEL/EPE NÃO entra no cálculo (a convenção de UF das usinas interestaduais muda
    ao longo da série deles); fica ao lado, só para comparação.
    """
    print("\n[1c] Reconstrução anual entre fotografias (âncoras) e validação leave-one-out")
    f = pd.read_csv(SAIDA / "fotografias_potencia_uf_fonte.csv")
    f = f[f.fonte_dado.str.contains("Wayback")]
    ancoras = [a for a in sorted(f.data_ref.unique())
               if a in ("2008-01-12", "2011-08-12", "2016-03-15", "2022-08-01", "2024-03-01",
                        "2024-09-18", "2025-01-06", "2025-05-21", "2025-09-22", "2026-01-16")]
    R = {a: f[(f.data_ref == a) & f.origem.isin(RENOVAVEIS)].groupby("uf").potencia_mw.sum().reindex(UFS).fillna(0)
         for a in ancoras}
    T = {a: f[f.data_ref == a].groupby("uf").potencia_mw.sum().reindex(UFS).fillna(0) for a in ancoras}
    # adições por unidade (UF principal pela CEG, como no SIGA)
    p = baixar(URL_LIBERACAO, BRUTO / "liberacao_detalhado.csv", minimo=1_000_000)
    d = pd.read_csv(p, sep=";", dtype=str, encoding="utf-8-sig")
    d["kw"] = pd.to_numeric(d.MdaPotenciaLiberadaComercial.str.replace(",", ".", regex=False), errors="coerce")
    d["dt"] = pd.to_datetime(d.DatLiberOpComerRealizado, errors="coerce")
    d["uf"] = d.CodCEG.str.extract(r"^[A-Z]{3}\.[A-Z]{2}\.([A-Z]{2})\.")[0]
    d = d[d.uf.isin(UFS) & d.dt.notna() & d.kw.notna()]

    def adic(a, b, so_renovavel):
        x = d[(d.dt > pd.Timestamp(a)) & (d.dt <= pd.Timestamp(b))]
        if so_renovavel:
            x = x[x.DscOrigemCombustivel.isin(RENOVAVEIS)]
        return (x.groupby("uf").kw.sum() / 1000).reindex(UFS).fillna(0)

    def ponte(X, so_ren, a1, a2, t):
        """X em t (entre as âncoras a1 e a2): fotografia a1 + adições até t + parte proporcional do fechamento."""
        t1, t2, tt = pd.Timestamp(a1), pd.Timestamp(a2), pd.Timestamp(t)
        fecha = X[a2] - (X[a1] + adic(a1, a2, so_ren))
        w = (tt - t1).days / max((t2 - t1).days, 1)
        return (X[a1] + adic(a1, t, so_ren) + fecha * w).clip(lower=0), fecha

    # leave-one-out: prever uma âncora do meio com as vizinhas
    loo = []
    for meio, a1, a2 in (("2011-08-12", "2008-01-12", "2016-03-15"), ("2016-03-15", "2011-08-12", "2022-08-01")):
        r, _ = ponte(R, True, a1, a2, meio)
        t, _ = ponte(T, False, a1, a2, meio)
        for uf in UFS:
            per_p = r[uf] / t[uf] * 100 if t[uf] else None
            per_r = R[meio][uf] / T[meio][uf] * 100 if T[meio][uf] else None
            loo.append({"ancora_prevista": meio, "ancoras_usadas": f"{a1} e {a2}", "uf": uf,
                        "renovavel_previsto_mw": round(r[uf], 1), "renovavel_real_mw": round(R[meio][uf], 1),
                        "total_previsto_mw": round(t[uf], 1), "total_real_mw": round(T[meio][uf], 1),
                        "per_previsto_pct": None if per_p is None else round(per_p, 2),
                        "per_real_pct": None if per_r is None else round(per_r, 2),
                        "erro_pp": None if per_p is None or per_r is None else round(per_p - per_r, 2)})
    loo = pd.DataFrame(loo)
    gravar(loo, "validacao_reconstrucao_leave_one_out.csv")
    mae = loo.groupby("ancora_prevista").erro_pp.apply(lambda x: x.abs().mean())
    print("   erro médio absoluto do PER previsto (pontos percentuais): " +
          ", ".join(f"{k}: {v:.2f}" for k, v in mae.items()) +
          f" | maior erro individual: {loo.erro_pp.abs().max():.1f} pp")

    t_aneel = capacidade_oficial()
    linhas = []
    for a1, a2 in zip(ancoras[:-1], ancoras[1:]):
        for ano in range(pd.Timestamp(a1).year, pd.Timestamp(a2).year + 1):
            tt = pd.Timestamp(f"{ano}-12-31")
            if not (pd.Timestamp(a1) <= tt <= pd.Timestamp(a2)):
                continue
            r, fr = ponte(R, True, a1, a2, tt)
            t, ft = ponte(T, False, a1, a2, tt)
            for uf in UFS:
                if not t[uf]:
                    continue
                aneel = t_aneel.loc[uf, ano] if ano in t_aneel.columns and not pd.isna(t_aneel.loc[uf, ano]) else None
                linhas.append({"uf": uf, "ano": ano, "valor": round(min(r[uf] / t[uf] * 100, 100), 2),
                               "unidade": "% da potência fiscalizada (reconstruído entre fotografias)",
                               "fonte": "reconstrução: fotografias ANEEL (BIG/SIGA) + liberações por unidade (ANEEL)",
                               "renovavel_mw": round(r[uf], 1), "total_mw": round(t[uf], 1),
                               "total_aneel_oficial_mw": None if aneel is None else round(float(aneel), 1),
                               "dif_total_vs_aneel_pct": None if aneel is None else round((t[uf] / float(aneel) - 1) * 100, 1),
                               "ancora_anterior": a1, "ancora_seguinte": a2,
                               "fechamento_renovavel_mw": round(fr[uf], 1), "fechamento_total_mw": round(ft[uf], 1)})
    out = pd.DataFrame(linhas).drop_duplicates(["uf", "ano"], keep="last").sort_values(["uf", "ano"])
    gravar(out, "per_potencia_reconstruida_uf_ano.csv")
    return out


# ============================== PARTE 2: ONS (energia) ======================================
FONTE_ONS = {"Hidráulica": "Hídrica", "Eólica": "Eólica", "Fotovoltaica": "Solar", "Biomassa": "Biomassa",
             "Resíduos Industriais": "Biomassa"}  # Suzano Maranhão = licor negro (biomassa no SIGA)
COLS_ONS = ["din_instante", "id_estado", "cod_modalidadeoperacao", "nom_tipousina", "nom_tipocombustivel",
            "nom_usina", "ceg", "val_geracao"]


def agregar_ons(parquet):
    d = pd.read_parquet(parquet, columns=COLS_ONS)
    d = d[d.id_estado.isin(UFS)].copy()
    d["v"] = pd.to_numeric(d.val_geracao, errors="coerce")
    d["ano"] = d.din_instante.dt.year
    d["mes"] = d.din_instante.dt.month
    d["horas"] = d.v.notna().astype(int)
    d["v"] = d.v.fillna(0.0)
    chave = ["ano", "mes", "id_estado", "cod_modalidadeoperacao", "nom_tipousina", "nom_tipocombustivel",
             "nom_usina", "ceg"]
    g = d.groupby(chave, dropna=False).agg(mwh=("v", "sum"), horas=("horas", "sum")).reset_index()
    return g


def parte_ons(epe=None):
    print("\n[2] ONS geracao-usina-2 (geração por usina, UF x combustível)")
    j = requests.get(CKAN_ONS + "geracao-usina-2", headers=HDR, timeout=60).json()["result"]
    urls = {r["name"].replace("Geracao_Usina_2-", ""): r["url"] for r in j["resources"]
            if r.get("format") == "PARQUET"}
    cache = BRUTO / "ons_agregado"
    cache.mkdir(parents=True, exist_ok=True)
    partes, hoje = [], pd.Timestamp.today()
    for chave, url in urls.items():
        dest = cache / f"{chave}.csv"
        ano = int(chave[:4])
        mes = int(chave[5:7]) if len(chave) > 4 else 12
        corrente = (ano, mes) >= (hoje.year, hoje.month)  # o mês em curso muda a cada dia
        if dest.exists() and not corrente:
            partes.append(pd.read_csv(dest))
            continue
        tmp = BRUTO / "ons_tmp.parquet"
        print("   baixando", chave)
        baixar(url, tmp, minimo=10000)
        g = agregar_ons(tmp)
        g.to_csv(dest, index=False)
        tmp.unlink()
        partes.append(g)
    d = pd.concat(partes, ignore_index=True)
    d["grupo"] = d.nom_tipocombustivel.map(FONTE_ONS).fillna("Fóssil")
    d["mmgd"] = d.cod_modalidadeoperacao.eq("Pequenas Usinas (MMGD)")
    d = d.rename(columns={"id_estado": "uf", "nom_usina": "usina", "nom_tipocombustivel": "combustivel",
                          "cod_modalidadeoperacao": "modalidade", "nom_tipousina": "tipo_usina"})
    # usina x ano
    pa = (d.groupby(["ano", "uf", "usina", "ceg", "modalidade", "tipo_usina", "combustivel", "grupo"],
                    dropna=False).agg(gwh=("mwh", lambda s: s.sum() / 1000), meses=("mes", "nunique"))
          .reset_index())
    pa["gwh"] = pa.gwh.round(3)
    gravar(pa, "ons_geracao_usina_ano.csv")
    # UF x combustível x ano
    ufa = (d.groupby(["ano", "uf", "combustivel", "grupo", "mmgd"]).agg(
        gwh=("mwh", lambda s: s.sum() / 1000), n_usinas=("usina", "nunique")).reset_index())
    ufa["gwh"] = ufa.gwh.round(3)
    gravar(ufa, "ons_geracao_uf_fonte_ano.csv")
    # PER de energia por UF e ano + indicadores de cobertura.
    # "valor" EXCLUI a MMGD (micro e minigeração distribuída): as linhas de MMGD só existem no ONS a partir
    # de 2023, então incluí-las quebraria a definição da série; o painel também não conta GD.
    of = capacidade_oficial()
    g_epe = {}
    if epe is not None:
        g_epe = {(r.uf, r.ano): r.valor for r in epe[epe.tabela == "2.5"].itertuples()}

    def linha(g, uf, ano):
        sm = g[~g.mmgd]
        tot_sm = sm.mwh.sum() / 1000
        ren_sm = sm[sm.grupo.isin(RENOVAVEIS)].mwh.sum() / 1000
        tot_cm = g.mwh.sum() / 1000
        ren_cm = g[g.grupo.isin(RENOVAVEIS)].mwh.sum() / 1000
        meses = int(sm.groupby("mes").mwh.sum().gt(0).sum())
        if uf == "AL":
            epe_uf = sum(g_epe[(u, ano)] for u in UFS) if all((u, ano) in g_epe for u in UFS) else None
            cap = None
        else:
            epe_uf = g_epe.get((uf, ano))
            cap = of.loc[uf, ano] if ano in of.columns else None
        return {
            "uf": uf, "ano": ano,
            "valor": round(ren_sm / tot_sm * 100, 2) if tot_sm > 0 else None,
            "valor_com_mmgd": round(ren_cm / tot_cm * 100, 2) if tot_cm > 0 else None,
            "unidade": "% da geração (GWh) monitorada pelo ONS, sem MMGD",
            "fonte": "ONS geracao-usina-2",
            "gwh_total": round(tot_sm, 1), "gwh_renovavel": round(ren_sm, 1),
            "gwh_mmgd_estimada": round(g[g.mmgd].mwh.sum() / 1000, 1),
            "n_usinas": int(sm[sm.mwh > 0].usina.nunique()), "meses_com_dado": meses,
            "parcial": bool(meses < 12),
            "n_ufs_com_dado": int(sm[sm.mwh > 0].uf.nunique()) if uf == "AL" else None,
            "fator_de_uso_implicito_pct": round(tot_sm / (cap * 8.76) * 100, 1) if cap else None,
            "gwh_epe_tabela_2_5": epe_uf,
            "cobertura_ons_sobre_epe_pct": round(tot_sm / epe_uf * 100, 1) if epe_uf and tot_sm > 0 else None}

    linhas = [linha(g, uf, int(ano)) for (ano, uf), g in d.groupby(["ano", "uf"]) if g.mwh.sum() > 0]
    linhas += [linha(g, "AL", int(ano)) for ano, g in d.groupby("ano")]
    per = pd.DataFrame(linhas).sort_values(["uf", "ano"])
    gravar(per, "per_energia_ons_uf_ano.csv")
    return per, d


# ============================== PARTE 3: CFURH (hídrica) ====================================
def parte_cfurh(ons_d=None):
    print("\n[3] CFURH/ANEEL: geração hídrica por UF")
    p = baixar(URL_CFURH, BRUTO / "cfurh-geracao.csv", minimo=100000)
    d = pd.read_csv(p, sep=";", decimal=",", dtype={"CodCEG": str})
    n0 = len(d)
    # VlrGeracaoTotal vem repetido por agente quando a usina tem mais de um sócio: tomar o máximo
    d = d.groupby(["CodCEG", "DatCompetencia"], as_index=False).agg(
        gwh=("VlrGeracaoTotal", "max"), ano=("AnoCompetencia", "first"), nome=("NomEmpreendimento", "first"))
    print(f"   {n0} linhas -> {len(d)} usina-mês após remover a repetição por agente")
    d["gwh"] = d.gwh / 1000
    d["uf"] = d.CodCEG.str[5:7]
    d["tipo"] = d.CodCEG.str[:3]
    d = d[d.uf.isin(UFS) & (d.ano < pd.Timestamp.today().year)]
    d["meses"] = 1
    g = d.groupby(["uf", "ano"]).agg(gwh=("gwh", "sum"), n_usinas=("CodCEG", "nunique"),
                                      usina_meses=("meses", "sum")).reset_index()
    g["gwh"] = g.gwh.round(1)
    g["unidade"] = "GWh"
    g["fonte"] = "ANEEL CFURH cfurh-geracao.csv (UHE e PCH; sem repetição por agente)"
    g = g.sort_values(["uf", "ano"])
    gravar(g, "cfurh_geracao_hidrica_uf_ano.csv")
    if ons_d is not None:
        o = ons_d[ons_d.grupo.eq("Hídrica") & ~ons_d.mmgd].groupby(["uf", "ano"]).mwh.sum().div(1000).reset_index()
        o.columns = ["uf", "ano", "ons_hidraulica_gwh"]
        c = g.merge(o, on=["uf", "ano"], how="outer").sort_values(["uf", "ano"])
        c["razao_ons_sobre_cfurh"] = (c.ons_hidraulica_gwh / c.gwh).round(3)
        # fora de +-10%: ou os universos diferem (ONS inclui PCH/CGH pequenas; CFURH só quem paga royalties)
        # ou há erro de um dos lados (CFURH tem Belo Monte jun/2019 = 14,3 TWh, o triplo do máximo físico)
        c["fora_de_10pct"] = c.razao_ons_sobre_cfurh.notna() & ((c.razao_ons_sobre_cfurh < 0.9) | (c.razao_ons_sobre_cfurh > 1.1))
        gravar(c[["uf", "ano", "gwh", "ons_hidraulica_gwh", "razao_ons_sobre_cfurh", "fora_de_10pct"]].rename(
            columns={"gwh": "cfurh_gwh"}), "cfurh_vs_ons_hidrica.csv")
    return g


# ============================== PARTE 4: geração distribuída solar ==========================
def parte_gd():
    print("\n[4] ANEEL: geração distribuída fotovoltaica (potência acumulada por ano de conexão)")
    p = baixar(URL_GD_FV, BRUTO / "gd_fv_tec.parquet", minimo=1_000_000, timeout=900)
    d = pd.read_parquet(p, columns=["CodGeracaoDistribuida", "DatConexao", "MdaPotenciaInstalada"])
    d["uf"] = d.CodGeracaoDistribuida.str[3:5]
    d = d[d.uf.isin(UFS)].copy()
    d["dt"] = pd.to_datetime(d.DatConexao, errors="coerce")
    sem_data = int(d.dt.isna().sum())
    d = d.dropna(subset=["dt"])
    d["ano"] = d.dt.dt.year
    por_ano = d.groupby(["uf", "ano"]).agg(adicao_mw=("MdaPotenciaInstalada", lambda s: s.sum() / 1000),
                                          n_sistemas=("CodGeracaoDistribuida", "count")).reset_index()
    linhas = []
    for uf, g in por_ano.groupby("uf"):
        acum = 0.0
        for r in g.sort_values("ano").itertuples():
            acum += r.adicao_mw
            linhas.append({"uf": uf, "ano": int(r.ano), "valor": round(acum, 1),
                           "adicao_no_ano_mw": round(r.adicao_mw, 1), "n_sistemas_no_ano": int(r.n_sistemas),
                           "unidade": "MW acumulados (fim do ano, sistemas hoje cadastrados)",
                           "fonte": "ANEEL empreendimento-gd-informacoes-tecnicas-fotovoltaica"})
    out = pd.DataFrame(linhas).sort_values(["uf", "ano"])
    print(f"   registros sem data de conexão ignorados: {sem_data}")
    gravar(out, "gd_solar_potencia_acumulada_uf_ano.csv")
    return out


# ============================== PARTE 5: testes de linha de base ============================
def parte_testes(per_pot=None, per_en=None):
    print("\n[5] Testes contra o painel e contra as linhas de base da ficha")
    resultados = []
    pot = pd.read_csv(SAIDA / "per_potencia_fotografias_uf.csv")
    # (a) o ponto da base do painel reproduz valores.csv?
    if VALORES_CSV.exists():
        v = pd.read_csv(VALORES_CSV, dtype=str)
        v = v[(v.codigo == CODIGO) & v.campo.isna() & v.ano.isna()].set_index("uf").valor.astype(float)
        base = pot[pot.fonte.str.contains("base do painel") & (pot.uf != "AL")].set_index("uf").valor
        for uf in UFS:
            resultados.append({"teste": "SIGA base do painel (2026-08-25) x valores.csv", "uf": uf,
                               "calculado": base.get(uf), "referencia": v.get(uf),
                               "diferenca": round(base.get(uf, float("nan")) - v.get(uf, float("nan")), 2)})
    al = pot[pot.fonte.str.contains("base do painel") & (pot.uf == "AL")].valor
    if len(al):
        resultados.append({"teste": "SIGA base do painel (2026-08-25), AL ponderado x painel (85,2)", "uf": "AL",
                           "calculado": float(al.iloc[0]), "referencia": 85.2,
                           "diferenca": round(float(al.iloc[0]) - 85.2, 2)})
    # (b) candidatos fechados às duas linhas de base (65,24 e 62,5), tolerância de 0,05 ponto
    cand = []
    for ref in ("2025-05-21", "2025-09-22"):
        g = pot[(pot.data_ref == ref) & pot.fonte.str.contains("SIGA via Wayback")]
        if len(g):
            cand.append((f"SIGA {ref}, AL ponderado pela potência", float(g[g.uf == "AL"].valor.iloc[0])))
            cand.append((f"SIGA {ref}, média simples das 9 UFs", float(g[g.uf != "AL"].valor.mean())))
    if per_en is not None:
        e = per_en[per_en.ano == 2025]
        if len(e):
            cand.append(("ONS 2025, AL ponderado pela geração, com MMGD", float(e[e.uf == "AL"].valor_com_mmgd.iloc[0])))
            cand.append(("ONS 2025, AL ponderado pela geração, sem MMGD", float(e[e.uf == "AL"].valor.iloc[0])))
            ufs = e[e.uf != "AL"]
            cand.append((f"ONS 2025, média simples das {ufs.valor_com_mmgd.notna().sum()} UFs com dado, com MMGD",
                         float(ufs.valor_com_mmgd.mean())))
            cand.append((f"ONS 2025, média simples das {ufs.valor.notna().sum()} UFs com dado, sem MMGD "
                         "(AC não tem geração monitorada fora da MMGD)", float(ufs.valor.mean())))
    for nome, val in cand:
        for alvo in (65.24, 62.5):
            resultados.append({"teste": f"{nome} x linha de base {alvo}", "uf": "AL", "calculado": round(val, 2),
                               "referencia": alvo, "diferenca": round(val - alvo, 2),
                               "reproduz_0_05": abs(val - alvo) <= 0.05})
    r = pd.DataFrame(resultados)
    gravar(r, "teste_linha_base.csv")
    print(r.to_string(index=False))
    return r


# ------------------------------------------------------------------------------------------
if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--so", choices=["siga", "big", "ons", "cfurh", "gd", "testes", "potencia", "epe", "hibrida"])
    a = ap.parse_args()
    BRUTO.mkdir(parents=True, exist_ok=True)
    per_pot = per_en = ons_d = epe = None
    if a.so in (None, "potencia", "siga", "big", "hibrida", "testes"):
        per_pot = parte_potencia()
    if a.so in (None, "epe", "hibrida", "ons", "cfurh", "testes"):
        epe = parte_epe()
    if a.so in (None, "hibrida", "testes"):
        parte_hibrida(epe)
    if a.so in (None, "ons", "cfurh", "testes"):
        per_en, ons_d = parte_ons(epe)
    if a.so in (None, "cfurh"):
        parte_cfurh(ons_d)
    if a.so in (None, "gd"):
        parte_gd()
    if a.so in (None, "testes"):
        parte_testes(per_pot, per_en)
