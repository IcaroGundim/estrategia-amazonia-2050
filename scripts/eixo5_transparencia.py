"""Eixo 5 / I5.5.2 — transparência pública dos governos estaduais.

A ficha aponta a Escala Brasil Transparente 360 (EBT 360, CGU, nota 0 a 10). O
Mapa Brasil Transparente (mbt.cgu.gov.br) segue fora do ar ("previsão de
retorno: novembro/2026") e o dadosabertos.cgu.gov.br não resolve, mas as páginas
de resultado das duas edições estão no Wayback Machine com a nota de cada estado
embutida. Não houve edição depois de 2020.

  EBT 360 1ª avaliação (2018)  .../escala_brasil_transparente/200000005
  EBT 360 2ª avaliação (2020)  .../escala_brasil_transparente/66

Valor principal = nota EBT 360 do governo do estado (0 a 10); série 2018, 2020.

A avaliação nacional que segue em curso é o PNTP — Programa Nacional de
Transparência Pública (Atricon e Tribunais de Contas), anual desde 2022, índice
de 0 a 100% e nível (Diamante ≥ 95%, Ouro 85-94%, Prata 75-84%, todos com 100%
dos critérios essenciais). Entra como campo auxiliar, junto com o ITGP da
Transparência Internacional (governança, 0 a 100), e as séries completas vão
para a tabela de detalhe. Trocar a fonte do valor principal é decisão de
método, não de coleta.

Arquivos locais em dados/eixo5/transparencia/ (fora do versionamento):
  avaliacoes_pntp_AAAA.xlsx  de radardatransparencia.atricon.org.br/dados/dados_pntp_AAAA.zip
  itgp2022.xlsx, itgp2025.xlsx  da Transparência Internacional Brasil

Uso:  python scripts/eixo5_transparencia.py [--aplicar]
"""
from __future__ import annotations

import csv
import re
import sys
import unicodedata
import urllib.request
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from estrategia_set2026 import UFS, aplica  # noqa: E402
from proposta import Proposta  # noqa: E402

RAIZ = Path(__file__).resolve().parents[1]
BRUTO = RAIZ / "dados" / "eixo5" / "transparencia"
CSV_PAINEL = RAIZ / "dashboard" / "public" / "data" / "csv"
NOMES = {"AC": "Acre", "AP": "Amapá", "AM": "Amazonas", "MA": "Maranhão", "MT": "Mato Grosso",
         "PA": "Pará", "RO": "Rondônia", "RR": "Roraima", "TO": "Tocantins"}
EBT = {
    "2018": "http://web.archive.org/web/20260213102422id_/https://mbt.cgu.gov.br/publico/avaliacao/escala_brasil_transparente/200000005",
    "2020": "http://web.archive.org/web/20260205211349id_/https://mbt.cgu.gov.br/publico/avaliacao/escala_brasil_transparente/66",
}
PNTP_URL = "https://radardatransparencia.atricon.org.br/dados/dados_pntp_{ano}.zip"
ITGP_URL = {"2022": "https://comunidade.transparenciainternacional.org.br/asset/204:itgp-executivo-estadual-base-de-dados",
            "2025": "https://comunidade.transparenciainternacional.org.br/asset/364:itgp-e-2025banco-de-dados"}


def sem_acento(texto: str) -> str:
    return unicodedata.normalize("NFD", str(texto)).encode("ascii", "ignore").decode().strip().upper()


UF_POR_NOME = {sem_acento(nome): uf for uf, nome in NOMES.items()}


def ebt(ano: str) -> dict[str, float]:
    copia = BRUTO / f"ebt360_{ano}.html"
    if not copia.exists():
        req = urllib.request.Request(EBT[ano], headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=120) as r:
            copia.write_bytes(r.read())
    pagina = copia.read_text(encoding="utf-8", errors="ignore")
    notas = {m.group(1).upper(): float(m.group(2)) for m in re.finditer(
        r'"uf":"(\w\w)","estado":"[^"]*","tooltip":"[^"]*?Nota do estado: ([\d.]+)', pagina)}
    return {uf: notas[uf] for uf in UFS}


def pntp(ano: str) -> dict[str, dict]:
    ws = openpyxl.load_workbook(BRUTO / f"avaliacoes_pntp_{ano}.xlsx", read_only=True).worksheets[0]
    linhas = ws.iter_rows(values_only=True)
    cab = {k: i for i, k in enumerate(next(linhas)) if k}
    col = lambda *nomes: next(cab[n] for n in nomes if n in cab)
    poder, uf_col = col("Poder", "Poder do Ente Público"), col("UF", "Estado")
    saida = {}
    for r in linhas:
        if r[poder] == "Executivo" and r[cab["Esfera"]] == "Estadual":
            uf = UF_POR_NOME.get(sem_acento(r[uf_col]))
            if uf:
                saida[uf] = {"indice": round(r[cab["Índice de Transparência"]] * 100, 2),
                             "essenciais": round(r[cab["% das Essenciais"]] * 100, 2),
                             "nivel": str(r[cab["Nível de Transparência"]]).strip()}
    return saida


def itgp(ano: str) -> dict[str, tuple[float, str]]:
    wb = openpyxl.load_workbook(BRUTO / f"itgp{ano}.xlsx", read_only=True, data_only=True)
    ws = wb.worksheets[0]
    saida = {}
    for r in ws.iter_rows(values_only=True):
        if ano == "2022" and r[1] and sem_acento(r[1]) in UF_POR_NOME:
            saida[UF_POR_NOME[sem_acento(r[1])]] = (round(float(r[3]), 1), str(r[4]).strip().capitalize())
        if ano == "2025" and r[1] and sem_acento(r[1]) in UF_POR_NOME:
            saida[UF_POR_NOME[sem_acento(r[1])]] = (round(float(r[3]), 1), str(r[4]).strip().capitalize())
    return saida


def main():
    notas = {ano: ebt(ano) for ano in EBT}
    avaliacoes = {ano: pntp(ano) for ano in ("2022", "2023", "2024", "2025")}
    itgps = {ano: itgp(ano) for ano in ITGP_URL}

    p = Proposta("eixo5_transparencia", fonte="CGU (EBT 360, via Wayback Machine); Atricon (PNTP); Transparência Internacional (ITGP)",
                 descricao="I5.5.2: nota EBT 360 2018 e 2020; PNTP 2025 e ITGP 2025 como auxiliares")
    for uf in UFS:
        for ano, n in notas.items():
            p.valor("I5.5.2", uf, n[uf], ano=int(ano), nota=f"EBT 360, {'1ª' if ano == '2018' else '2ª'} avaliação ({ano})")
        p.valor("I5.5.2", uf, notas["2020"][uf], nota="EBT 360, 2ª avaliação (2020), a última publicada pela CGU")
        a = avaliacoes["2025"][uf]
        p.valor("I5.5.2", uf, a["indice"], campo="pntpIndice2025")
        p.valor("I5.5.2", uf, a["nivel"], campo="pntpNivel2025")
        p.valor("I5.5.2", uf, a["essenciais"], campo="pntpEssenciais2025")
        p.valor("I5.5.2", uf, itgps["2025"][uf][0], campo="itgp2025")

    linhas = []
    for ano, n in notas.items():
        for uf in UFS:
            linhas.append([uf, "EBT 360 (CGU)", ano, n[uf], "0 a 10", "", "", EBT[ano]])
    for ano, dados in avaliacoes.items():
        for uf in UFS:
            if uf in dados:
                d = dados[uf]
                linhas.append([uf, "PNTP (Atricon)", ano, d["indice"], "0 a 100%", d["nivel"], d["essenciais"], PNTP_URL.format(ano=ano)])
    for ano, dados in itgps.items():
        for uf in UFS:
            if uf in dados:
                linhas.append([uf, "ITGP (Transparência Internacional)", ano, dados[uf][0], "0 a 100", dados[uf][1], "", ITGP_URL[ano]])
    with (CSV_PAINEL / "transparencia_estados.csv").open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(["uf", "avaliacao", "edicao", "nota", "escala", "nivel", "essenciais_pct", "fonte"])
        w.writerows(linhas)
    print("tabela: transparencia_estados.csv", len(linhas), "linhas")
    for ano, n in notas.items():
        print("EBT", ano, n, "| >= 9:", [uf for uf in UFS if n[uf] >= 9])
    for ano, d in avaliacoes.items():
        print("PNTP", ano, {uf: (d[uf]["indice"], d[uf]["nivel"]) for uf in d})
    print("ITGP", itgps)

    caminho = p.gravar()
    print("proposta:", caminho)
    if "--aplicar" in sys.argv and caminho:
        print("aplicadas:", aplica(caminho), "células")


if __name__ == "__main__":
    main()
